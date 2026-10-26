# scripts/build_company_docs.py
"""Build the NovaPay company corpus.

Each file in data/company_src/ is a short header (key: value lines), a blank
line, then the document body in light Markdown (# heading, ## heading, - bullet).
The header's `format` decides the single output file written to data/company/:
md, txt, docx or pdf. Title, author and date go into DOCX and PDF metadata.

    python -m scripts.build_company_docs
"""
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fpdf import FPDF

SRC_DIR = Path("data/company_src")
OUT_DIR = Path("data/company")
FORMATS = ("md", "txt", "docx", "pdf")


def parse_source(text: str) -> tuple[dict, str]:
    """Split a source file into its header dict and its body."""
    header, sep, body = text.partition("\n\n")
    if not sep:
        raise ValueError("no blank line between the header and the body")
    meta = {}
    for line in header.splitlines():
        key, colon, value = line.partition(":")
        if not colon:
            raise ValueError(f"header line without ':': {line!r}")
        meta[key.strip().lower()] = value.strip()
    for key in ("title", "format"):
        if not meta.get(key):
            raise ValueError(f"header is missing '{key}'")
    if meta["format"] not in FORMATS:
        raise ValueError(f"format must be one of {FORMATS}, got {meta['format']!r}")
    return meta, body.strip() + "\n"


def line_kind(line: str) -> tuple[str, str]:
    """Classify one body line as h1, h2, bullet or para, without its marker."""
    if line.startswith("# "):
        return "h1", line[2:]
    if line.startswith("## "):
        return "h2", line[3:]
    if line.startswith("- "):
        return "bullet", line[2:]
    return "para", line


def parse_date(value: str) -> datetime:
    return datetime.strptime(value, "%Y-%m-%d").replace(tzinfo=timezone.utc)


def write_text(meta: dict, body: str, path: Path) -> None:
    path.write_text(body, encoding="utf-8")


def write_docx(meta: dict, body: str, path: Path) -> None:
    doc = Document()
    for line in body.splitlines():
        if not line.strip():
            continue
        kind, text = line_kind(line)
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    props = doc.core_properties
    props.title = meta["title"]
    props.author = meta.get("author", "")
    if meta.get("date"):
        props.created = parse_date(meta["date"])
    doc.save(path)


def write_pdf(meta: dict, body: str, path: Path) -> None:
    # The built-in Helvetica font only covers Latin-1 characters.
    for ch in body + meta["title"]:
        if ord(ch) > 255:
            raise ValueError(f"character {ch!r} cannot be written to a PDF with the built-in font; "
                             "replace it (for example an em dash with '-')")
    pdf = FPDF()
    pdf.set_title(meta["title"])
    if meta.get("author"):
        pdf.set_author(meta["author"])
    if meta.get("date"):
        pdf.set_creation_date(parse_date(meta["date"]))
    pdf.set_margins(20, 20, 20)
    pdf.add_page()
    for line in body.splitlines():
        if not line.strip():
            pdf.ln(3)
            continue
        kind, text = line_kind(line)
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 8, text, new_x="LMARGIN", new_y="NEXT")
        elif kind == "h2":
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 11)
            if kind == "bullet":
                text = "- " + text
            pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


WRITERS = {"md": write_text, "txt": write_text, "docx": write_docx, "pdf": write_pdf}


def main() -> None:
    sources = sorted(SRC_DIR.glob("*.md"))
    if not sources:
        sys.exit(f"No source files found in {SRC_DIR}/ (run from the repo root)")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for old in OUT_DIR.iterdir():  # remove last run's output so renamed sources leave nothing behind
        if old.suffix.lstrip(".") in FORMATS:
            old.unlink()

    counts = Counter()
    for src in sources:
        try:
            meta, body = parse_source(src.read_text(encoding="utf-8"))
            out = OUT_DIR / f"{src.stem}.{meta['format']}"
            WRITERS[meta["format"]](meta, body, out)
        except ValueError as exc:
            sys.exit(f"{src}: {exc}")
        counts[meta["format"]] += 1
        print(f"{out}  ({len(body.split())} words)")

    print(f"Done: {len(sources)} documents ->", ", ".join(f"{n} {fmt}" for fmt, n in sorted(counts.items())))


if __name__ == "__main__":
    main()
