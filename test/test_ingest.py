# tests/test_ingest.py
import io
import os

import docx
import psycopg
import pytest
from fpdf import FPDF

from app.ingest import IngestError, content_hash, load_document, save_document

# Never Neon: defaults to the local Docker database, override with TEST_DATABASE_URL.
TEST_DB = os.getenv("TEST_DATABASE_URL", "postgresql://postgres:postgres@localhost:5433/investigation")


def test_whitespace_differences_give_the_same_hash():
    assert content_hash("Payments  failed.\n\nAt 9:00") == content_hash(" Payments failed. At 9:00 ")


def test_different_text_gives_a_different_hash():
    assert content_hash("refund") != content_hash("refunds")


def test_markdown_title_comes_from_first_heading():
    doc = load_document("runbook.md", b"# Refund runbook\r\n\r\n\r\nStep one.\n")
    assert doc.title == "Refund runbook"
    assert doc.text == "# Refund runbook\n\nStep one."
    assert doc.metadata == {"file_name": "runbook.md", "format": "md"}


def test_docx_keeps_headings_and_author():
    d = docx.Document()
    d.core_properties.author = "Ana"
    d.add_heading("Incident 42", level=1)
    d.add_paragraph("Payments failed for 20 minutes.")
    buf = io.BytesIO()
    d.save(buf)

    doc = load_document("incident.docx", buf.getvalue())
    assert doc.title == "Incident 42"
    assert doc.text == "# Incident 42\n\nPayments failed for 20 minutes."
    assert doc.metadata["author"] == "Ana"


def test_pdf_text_page_count_and_fallback_title():
    pdf = FPDF()
    pdf.set_author("Ana")
    pdf.set_font("Helvetica", size=12)
    for line in ("First page", "Second page"):
        pdf.add_page()
        pdf.cell(text=line)

    doc = load_document("march_report.pdf", bytes(pdf.output()))
    assert "First page" in doc.text and "Second page" in doc.text
    assert doc.metadata["page_count"] == 2
    assert doc.metadata["author"] == "Ana"
    assert doc.title == "march report"  # no title inside the PDF, so the file name is used


def test_rejects_unsupported_broken_and_empty_files():
    with pytest.raises(IngestError, match="unsupported"):
        load_document("photo.png", b"...")
    with pytest.raises(IngestError, match="not a valid DOCX"):
        load_document("broken.docx", b"not a zip")
    with pytest.raises(IngestError, match="no text"):
        load_document("empty.txt", b"  \n\n ")


@pytest.fixture
def conn():
    try:
        connection = psycopg.connect(TEST_DB, connect_timeout=3)
    except psycopg.OperationalError:
        pytest.skip("local Postgres is not running (docker compose up -d)")
    with connection:
        with connection.transaction(force_rollback=True):  # the test leaves no rows behind
            yield connection


def test_add_then_duplicate_then_update(conn):
    v1 = load_document("pytest_note.md", b"# Note\n\nThe refund job runs at 02:00.")
    assert save_document(conn, v1, source="upload")[0] == "added"
    assert save_document(conn, v1, source="upload")[0] == "duplicate"

    v2 = load_document("pytest_note.md", b"# Note\n\nThe refund job runs at 03:00.")
    status, doc_id = save_document(conn, v2, source="upload")
    assert status == "updated"
    version = conn.execute("SELECT version FROM documents WHERE id = %s", (doc_id,)).fetchone()[0]
    assert version == 2